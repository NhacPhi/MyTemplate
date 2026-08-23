import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('Tool/data/Localizations.xlsx', data_only=True)
found = False
for sheet in wb.sheetnames:
    ws = wb[sheet]
    for r in ws.iter_rows(values_only=True):
        if r[0] == 'STR_COOLDOWN_REDUCED':
            print(f"Found in sheet {sheet}: {r}")
            found = True

if not found:
    print("Not found in Localizations.xlsx")
