import openpyxl
import time
import sys

sys.stdout.reconfigure(encoding='utf-8')

shop_path = 'Tool/data/Shop.xlsx'

armor_products = [
    # Set 01 - Kim Vũ Long Lân Khải (Legendary Đấu Sĩ Flagship)
    ('SHOP_ARMOR_01_A', 'ComboPack', 'JadeShop', 'Bundle', 'BUNDLE_ARMOR01_01', 1, 'Jade', 800, 1200, 1, 'Weekly', None, None, True, 1),
    ('SHOP_ARMOR_01_B', 'ComboPack', 'JadeShop', 'Bundle', 'BUNDLE_ARMOR01_02', 1, 'Jade', 800, 1200, 1, 'Weekly', None, None, True, 2),

    # Set 07 - Huyết Chiến Vương Giáp (Legendary Sát Thủ)
    ('SHOP_ARMOR_07_A', 'ComboPack', 'JadeShop', 'Bundle', 'BUNDLE_ARMOR07_01', 1, 'Jade', 800, 1200, 1, 'Weekly', None, None, True, 3),
    ('SHOP_ARMOR_07_B', 'ComboPack', 'JadeShop', 'Bundle', 'BUNDLE_ARMOR07_02', 1, 'Jade', 800, 1200, 1, 'Weekly', None, None, True, 4),

    # Set 08 - Thanh Long Bố Lân Bào (Legendary Đấu Sĩ Xuyên Giáp)
    ('SHOP_ARMOR_08_A', 'ComboPack', 'JadeShop', 'Bundle', 'BUNDLE_ARMOR08_01', 1, 'Jade', 800, 1200, 1, 'Weekly', None, None, True, 5),
    ('SHOP_ARMOR_08_B', 'ComboPack', 'JadeShop', 'Bundle', 'BUNDLE_ARMOR08_02', 1, 'Jade', 800, 1200, 1, 'Weekly', None, None, True, 6),
]

armor_bundle_contents = [
    # Set 01
    ('BUNDLE_ARMOR01_01', 'Armor01_Helmet', 1),
    ('BUNDLE_ARMOR01_01', 'Armor01_Chestplate', 1),
    ('BUNDLE_ARMOR01_01', 'Armor01_Gloves', 1),
    ('BUNDLE_ARMOR01_02', 'Armor01_Boots', 1),
    ('BUNDLE_ARMOR01_02', 'Armor01_Belt', 1),
    ('BUNDLE_ARMOR01_02', 'Armor01_Ring', 1),

    # Set 07
    ('BUNDLE_ARMOR07_01', 'Armor07_Helmet', 1),
    ('BUNDLE_ARMOR07_01', 'Armor07_Chestplate', 1),
    ('BUNDLE_ARMOR07_01', 'Armor07_Gloves', 1),
    ('BUNDLE_ARMOR07_02', 'Armor07_Boots', 1),
    ('BUNDLE_ARMOR07_02', 'Armor07_Belt', 1),
    ('BUNDLE_ARMOR07_02', 'Armor07_Ring', 1),

    # Set 08
    ('BUNDLE_ARMOR08_01', 'Armor08_Helmet', 1),
    ('BUNDLE_ARMOR08_01', 'Armor08_Chestplate', 1),
    ('BUNDLE_ARMOR08_01', 'Armor08_Gloves', 1),
    ('BUNDLE_ARMOR08_02', 'Armor08_Boots', 1),
    ('BUNDLE_ARMOR08_02', 'Armor08_Belt', 1),
    ('BUNDLE_ARMOR08_02', 'Armor08_Ring', 1),
]

for attempt in range(5):
    try:
        wb = openpyxl.load_workbook(shop_path)

        # 1. Update ShopProducts
        ws_prod = wb['ShopProducts']
        existing_rows = []
        for r in ws_prod.iter_rows(values_only=True):
            if r[0] and not str(r[0]).startswith('SHOP_ARMOR_'):
                existing_rows.append(r)
        
        ws_prod.delete_rows(1, ws_prod.max_row)
        for r in existing_rows:
            ws_prod.append(list(r))
        for r in armor_products:
            ws_prod.append(list(r))

        # 2. Update BundleContents
        ws_bundle = wb['BundleContents']
        existing_bundles = []
        for r in ws_bundle.iter_rows(values_only=True):
            if r[0] and not str(r[0]).startswith('BUNDLE_ARMOR'):
                existing_bundles.append(r)
        
        ws_bundle.delete_rows(1, ws_bundle.max_row)
        for r in existing_bundles:
            ws_bundle.append(list(r))
        for r in armor_bundle_contents:
            ws_bundle.append(list(r))

        wb.save(shop_path)
        print("Successfully saved Shop.xlsx with Set 01, Set 07, and Set 08 Armor Bundles!")
        break
    except Exception as e:
        print(f"Attempt {attempt+1} failed: {e}. Retrying...")
        time.sleep(1)
