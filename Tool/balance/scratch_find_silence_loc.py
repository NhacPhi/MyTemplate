import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('Tool/data/Localizations.xlsx')
for s in wb.sheetnames:
    ws = wb[s]
    headers = [cell.value for cell in ws[1]]
    for r in range(2, ws.max_row + 1):
        k = str(ws.cell(r, 1).value).strip() if ws.cell(r, 1).value else ''
        if 'SILENCE' in k or 'Silence' in k or 'Câm' in str([ws.cell(r, c).value for c in range(1, len(headers)+1)]):
            print(f"Sheet: {s}, Row {r}: Headers={headers}")
            print(f"Values: {[ws.cell(r, c).value for c in range(1, len(headers)+1)]}")
