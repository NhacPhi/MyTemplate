import openpyxl, time, sys
sys.stdout.reconfigure(encoding='utf-8')

print("=== STARTING EXCEL UPDATE ===")

# 1. Update Localizations.xlsx
loc_saved = False
for attempt in range(10):
    try:
        wb_loc = openpyxl.load_workbook('Tool/data/Localizations.xlsx')
        ws_loc = wb_loc['Effect'] if 'Effect' in wb_loc.sheetnames else wb_loc.active
        
        existing_keys = {row[0].value for row in ws_loc.iter_rows(min_row=2) if row[0].value}
        
        if 'EFF_SILENCE_NAME' not in existing_keys:
            ws_loc.append(['EFF_SILENCE_NAME', 'Câm Lặng', 'Silence', 'Câm Lặng'])
            print("Added EFF_SILENCE_NAME to Localizations.xlsx")
            
        if 'EFF_SILENCE_DES' not in existing_keys:
            ws_loc.append(['EFF_SILENCE_DES', 'Không thể sử dụng Tuyệt kỹ và Kỹ năng', 'Cannot use Major and Ultimate skills', 'Không thể sử dụng Tuyệt kỹ và Kỹ năng'])
            print("Added EFF_SILENCE_DES to Localizations.xlsx")
            
        wb_loc.save('Tool/data/Localizations.xlsx')
        loc_saved = True
        print("Successfully saved Localizations.xlsx!")
        break
    except PermissionError:
        print(f"Localizations.xlsx locked (attempt {attempt+1}/10), retrying in 1s...")
        time.sleep(1)

# 2. Update GameConfig.xlsx
gc_saved = False
for attempt in range(10):
    try:
        wb = openpyxl.load_workbook('Tool/data/GameConfig.xlsx')
        
        # A. EffectConfig Sheet
        ws_eff = wb['EffectConfig']
        eff_keys = {row[0].value for row in ws_eff.iter_rows(min_row=2) if row[0].value}
        if 'EFF_Silence_def' not in eff_keys:
            ws_eff.append(['EFF_Silence_def', 'EFF_SILENCE_NAME', 'EFF_SILENCE_DES', 'Silence', 'None', 'None', 1, 100, 1])
            print("Added EFF_Silence_def to EffectConfig sheet in GameConfig.xlsx")
            
        # B. SkillConfig Sheet
        ws_sk = wb['SkillConfig']
        sk_headers = [c.value for c in ws_sk[1]]
        # Find column indices
        col_type = sk_headers.index('Type') if 'Type' in sk_headers else 3
        col_dmg = sk_headers.index('DamageMultiplier') if 'DamageMultiplier' in sk_headers else 4
        col_cd = sk_headers.index('MaxCooldown') if 'MaxCooldown' in sk_headers else 5
        col_stype = sk_headers.index('SkillType') if 'SkillType' in sk_headers else 6
        col_ttype = sk_headers.index('TargetType') if 'TargetType' in sk_headers else 7
        col_eff = sk_headers.index('Effect') if 'Effect' in sk_headers else 8
        
        for row in ws_sk.iter_rows(min_row=2):
            s_id = str(row[0].value).strip() if row[0].value else ''
            if s_id == 'GoldfishDemon_B':
                row[col_dmg].value = '0.8, 1.0, 1.2'
                row[col_cd].value = '0, 0, 0'
                row[col_stype].value = 'BasicAttack'
                row[col_ttype].value = 'SingleEnemy'
                row[col_eff].value = 'None'
                print(f"Updated SkillConfig: {s_id}")
            elif s_id == 'GoldfishDemon_M':
                row[col_type].value = 'MajorAttack'
                row[col_dmg].value = '0.7, 0.85, 1.0'
                row[col_cd].value = '3, 3, 3'
                row[col_stype].value = 'ActiveSkill'
                row[col_ttype].value = 'EnemyRow'
                row[col_eff].value = 'EFF_Silence_def'
                print(f"Updated SkillConfig: {s_id}")
            elif s_id == 'GoldfishDemon_U':
                row[col_type].value = 'EmpowerAttack'
                row[col_dmg].value = '1.8, 2.0, 2.3'
                row[col_cd].value = '4, 4, 4'
                row[col_stype].value = 'ActiveSkill'
                row[col_ttype].value = 'SingleEnemy'
                row[col_eff].value = 'EFF_Stun_def'
                print(f"Updated SkillConfig: {s_id}")
                
        # C. CharacterStat Sheet
        ws_stat = wb['CharacterStat']
        stat_headers = [c.value for c in ws_stat[1]]
        col_hp = stat_headers.index('hp') if 'hp' in stat_headers else 5
        col_atk = stat_headers.index('atk') if 'atk' in stat_headers else 6
        col_def = stat_headers.index('def') if 'def' in stat_headers else 7
        col_spd = stat_headers.index('speed') if 'speed' in stat_headers else 8
        col_shred = stat_headers.index('def_shred') if 'def_shred' in stat_headers else 9
        col_cr = stat_headers.index('crit_rare') if 'crit_rare' in stat_headers else 10
        col_cdmg = stat_headers.index('crit_dmg') if 'crit_dmg' in stat_headers else 11
        col_pen = stat_headers.index('penetration') if 'penetration' in stat_headers else 12
        col_res = stat_headers.index('crit_dmg_res') if 'crit_dmg_res' in stat_headers else 13
        
        for row in ws_stat.iter_rows(min_row=2):
            c_id = str(row[0].value).strip() if row[0].value else ''
            if c_id == 'GoldfishDemon':
                row[col_hp].value = 11364
                row[col_atk].value = 3350
                row[col_def].value = 530
                row[col_spd].value = 94
                row[col_shred].value = 14
                row[col_cr].value = 60
                row[col_cdmg].value = 100
                row[col_pen].value = 20
                row[col_res].value = 25
                print(f"Updated CharacterStat: {c_id}")
                
        # D. CharacterUpgrade Sheet
        ws_up = wb['CharacterUpgrade']
        for row in ws_up.iter_rows(min_row=2):
            c_id = str(row[0].value).strip() if row[0].value else ''
            if c_id == 'GoldfishDemon':
                row[1].value = 568
                row[2].value = 168
                row[3].value = 26
                print(f"Updated CharacterUpgrade: {c_id}")
                
        wb.save('Tool/data/GameConfig.xlsx')
        gc_saved = True
        print("Successfully saved GameConfig.xlsx!")
        break
    except PermissionError:
        print(f"GameConfig.xlsx locked (attempt {attempt+1}/10), retrying in 1s...")
        time.sleep(1)

# Verification
print("\n=== VERIFICATION OF SAVED EXCEL FILES ===")
wb_check = openpyxl.load_workbook('Tool/data/GameConfig.xlsx', data_only=True)

print("\n1. EffectConfig Sheet verification:")
for row in wb_check['EffectConfig'].iter_rows(values_only=True):
    if row[0] in ['EFF_Silence_def', 'EFF_Stun_def']:
        print(row)

print("\n2. SkillConfig Sheet verification:")
for row in wb_check['SkillConfig'].iter_rows(values_only=True):
    if row[0] and 'Goldfish' in str(row[0]):
        print(row[:9])

print("\n3. CharacterStat Sheet verification:")
for row in wb_check['CharacterStat'].iter_rows(values_only=True):
    if row[0] == 'GoldfishDemon':
        print(row[:14])
