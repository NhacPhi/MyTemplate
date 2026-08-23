import openpyxl
import time
import sys

sys.stdout.reconfigure(encoding='utf-8')

shop_path = 'Tool/data/Shop.xlsx'

# 4 Bundles for Set 07 & Set 08:
new_products = [
    ('SHOP_ARMOR_07_A', 'ComboPack', 'JadeShop', 'Bundle', 'BUNDLE_ARMOR07_01', 1, 'Jade', 800, 1200, 1, 'Weekly', None, None, True, 1),
    ('SHOP_ARMOR_07_B', 'ComboPack', 'JadeShop', 'Bundle', 'BUNDLE_ARMOR07_02', 1, 'Jade', 800, 1200, 1, 'Weekly', None, None, True, 2),
    ('SHOP_ARMOR_08_A', 'ComboPack', 'JadeShop', 'Bundle', 'BUNDLE_ARMOR08_01', 1, 'Jade', 800, 1200, 1, 'Weekly', None, None, True, 3),
    ('SHOP_ARMOR_08_B', 'ComboPack', 'JadeShop', 'Bundle', 'BUNDLE_ARMOR08_02', 1, 'Jade', 800, 1200, 1, 'Weekly', None, None, True, 4),
]

new_bundle_contents = [
    # Set 07 - Huyết Chiến Vương Giáp (Phần 1: Nón, Áo, Tay)
    ('BUNDLE_ARMOR07_01', 'Armor07_Helmet', 1),
    ('BUNDLE_ARMOR07_01', 'Armor07_Chestplate', 1),
    ('BUNDLE_ARMOR07_01', 'Armor07_Gloves', 1),
    
    # Set 07 - Huyết Chiến Vương Giáp (Phần 2: Giày, Thắt Lưng, Nhẫn)
    ('BUNDLE_ARMOR07_02', 'Armor07_Boots', 1),
    ('BUNDLE_ARMOR07_02', 'Armor07_Belt', 1),
    ('BUNDLE_ARMOR07_02', 'Armor07_Ring', 1),

    # Set 08 - Thanh Long Bố Lân Bào (Phần 1: Nón, Áo, Tay)
    ('BUNDLE_ARMOR08_01', 'Armor08_Helmet', 1),
    ('BUNDLE_ARMOR08_01', 'Armor08_Chestplate', 1),
    ('BUNDLE_ARMOR08_01', 'Armor08_Gloves', 1),

    # Set 08 - Thanh Long Bố Lân Bào (Phần 2: Giày, Thắt Lưng, Nhẫn)
    ('BUNDLE_ARMOR08_02', 'Armor08_Boots', 1),
    ('BUNDLE_ARMOR08_02', 'Armor08_Belt', 1),
    ('BUNDLE_ARMOR08_02', 'Armor08_Ring', 1),
]

for attempt in range(5):
    try:
        wb = openpyxl.load_workbook(shop_path)

        # 1. Update ShopProducts
        ws_prod = wb['ShopProducts']
        # Remove any previous armor bundle products
        existing_rows = []
        for r in ws_prod.iter_rows(values_only=True):
            if r[0] not in ['SHOP_ARMOR_07_A', 'SHOP_ARMOR_07_B', 'SHOP_ARMOR_08_A', 'SHOP_ARMOR_08_B']:
                existing_rows.append(r)
        
        ws_prod.delete_rows(1, ws_prod.max_row)
        for r in existing_rows:
            ws_prod.append(list(r))
        for r in new_products:
            ws_prod.append(list(r))

        # 2. Update BundleContents
        ws_bundle = wb['BundleContents']
        existing_bundles = []
        for r in ws_bundle.iter_rows(values_only=True):
            if r[0] not in ['BUNDLE_ARMOR07_01', 'BUNDLE_ARMOR07_02', 'BUNDLE_ARMOR08_01', 'BUNDLE_ARMOR08_02']:
                existing_bundles.append(r)
        
        ws_bundle.delete_rows(1, ws_bundle.max_row)
        for r in existing_bundles:
            ws_bundle.append(list(r))
        for r in new_bundle_contents:
            ws_bundle.append(list(r))

        wb.save(shop_path)
        print("Successfully saved Shop.xlsx with 4 Armor Bundles!")
        break
    except Exception as e:
        print(f"Attempt {attempt+1} failed: {e}. Retrying...")
        time.sleep(1)
