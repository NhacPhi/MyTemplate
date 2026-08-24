import openpyxl, time, sys
sys.stdout.reconfigure(encoding='utf-8')

for attempt in range(10):
    try:
        wb = openpyxl.load_workbook('Tool/data/Localizations.xlsx')
        ws = wb['Effect']
        
        headers = [c.value for c in ws[1]]
        print("Headers in Effect:", headers)
        col_en = headers.index('ENGLISH') + 1
        col_vn = headers.index('VIETNAMESE') + 1
        
        for r in range(2, ws.max_row + 1):
            k = str(ws.cell(r, 1).value).strip() if ws.cell(r, 1).value else ''
            if k == 'EFF_SILENCE_NAME':
                ws.cell(r, col_en).value = 'Silence'
                ws.cell(r, col_vn).value = 'Câm Lặng'
                # Clear extra columns if any
                for c in range(max(col_en, col_vn) + 1, ws.max_column + 1):
                    ws.cell(r, c).value = None
                print(f"Fixed EFF_SILENCE_NAME: EN='{ws.cell(r, col_en).value}', VN='{ws.cell(r, col_vn).value}'")
            elif k == 'EFF_SILENCE_DES':
                ws.cell(r, col_en).value = 'Cannot use Major and Ultimate skills'
                ws.cell(r, col_vn).value = 'Không thể sử dụng Tuyệt kỹ và Kỹ năng'
                for c in range(max(col_en, col_vn) + 1, ws.max_column + 1):
                    ws.cell(r, c).value = None
                print(f"Fixed EFF_SILENCE_DES: EN='{ws.cell(r, col_en).value}', VN='{ws.cell(r, col_vn).value}'")
                
        wb.save('Tool/data/Localizations.xlsx')
        print("Successfully saved fixed Localizations.xlsx!")
        break
    except PermissionError:
        print(f"Locked attempt {attempt+1}, retrying in 1s...")
        time.sleep(1)
