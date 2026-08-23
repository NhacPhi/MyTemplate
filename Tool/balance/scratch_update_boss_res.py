import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('Tool/data/GameConfig.xlsx')
ws = wb['CharacterStat']
ws_char = wb['Character']

char_class_map = {}
char_role_map = {}
for r in ws_char.iter_rows(min_row=2, values_only=True):
    if r[0]:
        char_class_map[r[0]] = r[4]
        char_role_map[r[0]] = r[3]

header = [cell.value for cell in ws[1]]
id_idx = header.index('ID')
cr_idx = header.index('crit_rare') if 'crit_rare' in header else header.index('crit_rate')
res_idx = header.index('crit_dmg_res')

for row in ws.iter_rows(min_row=2):
    cid = row[id_idx].value
    cls = char_class_map.get(cid, 'Character')
    role = char_role_map.get(cid, 'Fighter')
    
    if cls == 'Boss':
        row[cr_idx].value = 50
        if role == 'Tanker':
            row[res_idx].value = 30
        else:
            row[res_idx].value = 25
        print(f"Updated Boss {cid} ({role}): crit_rate = {row[cr_idx].value}%, crit_dmg_res = {row[res_idx].value}%")

wb.save('Tool/data/GameConfig.xlsx')
print("Successfully saved CharacterStat in GameConfig.xlsx")
