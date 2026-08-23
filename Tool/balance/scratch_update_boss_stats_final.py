import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('Tool/data/GameConfig.xlsx')
ws_char = wb['Character']
ws_stat = wb['CharacterStat']

# Get char class map
char_class_map = {}
for r in ws_char.iter_rows(min_row=2, values_only=True):
    if r[0]:
        char_class_map[r[0]] = r[4]

header = [cell.value for cell in ws_stat[1]]
print("CharacterStat Header:", header)

id_idx = header.index('ID')
cr_idx = header.index('crit_rare') if 'crit_rare' in header else header.index('crit_rate')
cd_idx = header.index('crit_dmg')
pen_idx = header.index('penetration')
res_idx = header.index('crit_dmg_res')

for row in ws_stat.iter_rows(min_row=2):
    cid = row[id_idx].value
    cls = char_class_map.get(cid, 'Character')
    
    if cls == 'Boss':
        row[cr_idx].value = 60    # 60% Crit Rate
        row[cd_idx].value = 100   # 100% Bonus Crit DMG (150% Base + 100% = 250% Total Crit DMG)
        row[pen_idx].value = 20   # 20% Penetration
        row[res_idx].value = 25   # 25% Crit DMG Res
        print(f"Updated Boss {cid}: CritRate={row[cr_idx].value}%, CritDMG={row[cd_idx].value} (Total 250%), Pen={row[pen_idx].value}%, CritRes={row[res_idx].value}%")

wb.save('Tool/data/GameConfig.xlsx')
print("Successfully saved all Boss stats in Tool/data/GameConfig.xlsx")
