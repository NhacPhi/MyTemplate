import openpyxl, time, sys
sys.stdout.reconfigure(encoding='utf-8')

# Ensure GameConfig.xlsx StageEnemies is cleanly written
exact_levels = [
    # Tutorial
    ('Tutorial_Default', 1, 'VanguardTiger', 5, False),
    ('Tutorial_Default', 2, 'VanguardTiger', 5, False),
    ('Tutorial_Default', 3, 'VanguardTiger', 5, False),
    
    # Region 1: Dong Hai
    ('Battle_GoldfishDemon', 1, 'Benborba', 10, False),
    ('Battle_GoldfishDemon', 2, 'Baborben', 10, False),
    ('Battle_GoldfishDemon', 3, 'Benborba', 10, False),
    ('Battle_GoldfishDemon', 5, 'GoldfishDemon', 12, True),
    
    ('Battle_ThirdDragonPrince', 1, 'CrabSolider', 15, False),
    ('Battle_ThirdDragonPrince', 2, 'SquidSolider', 15, False),
    ('Battle_ThirdDragonPrince', 3, 'CrabSolider', 15, False),
    ('Battle_ThirdDragonPrince', 4, 'GoldfishDemon', 20, False),
    ('Battle_ThirdDragonPrince', 5, 'ThirdDragonPrince', 20, True),
    
    # Region 2: Thien Dinh
    ('Battle_LiJing_Boss', 1, 'MarshalTianpeng', 20, False),
    ('Battle_LiJing_Boss', 2, 'LiJing_Boss', 25, True),
    ('Battle_LiJing_Boss', 3, 'HeavenlySoddier', 20, False),
    ('Battle_LiJing_Boss', 4, 'HeavenlySoddier', 20, False),
    ('Battle_LiJing_Boss', 5, 'ThirdPrinceNezha_Boss', 20, False),
    ('Battle_LiJing_Boss', 6, 'HeavenlySoddier', 20, False),
    
    ('Battle_ErlangShen_Boss', 1, 'HeavenlySoddier', 20, False),
    ('Battle_ErlangShen_Boss', 2, 'HeavenlySoddier', 20, False),
    ('Battle_ErlangShen_Boss', 3, 'HeavenlySoddier', 20, False),
    ('Battle_ErlangShen_Boss', 5, 'ErlangShen_Boss', 25, True),
    
    # Region 3: Cao Lao Trang & Hoang Phong
    ('Battle_YellowWindMonster', 1, 'VanguardTiger', 20, False),
    ('Battle_YellowWindMonster', 2, 'RatMonster', 20, False),
    ('Battle_YellowWindMonster', 3, 'VanguardTiger', 20, False),
    ('Battle_YellowWindMonster', 5, 'YellowWindMonster', 28, True),
    
    ('Battle_GoldHornKing', 1, 'Minions', 25, False),
    ('Battle_GoldHornKing', 2, 'Minions', 25, False),
    ('Battle_GoldHornKing', 3, 'Minions', 25, False),
    ('Battle_GoldHornKing', 4, 'SilverHornKing', 26, False),
    ('Battle_GoldHornKing', 5, 'GoldHornKing', 28, True),
    
    # Region 4: Hoa Diem Son
    ('Battle_TheBoySage', 1, 'YoungBufflalo', 25, False),
    ('Battle_TheBoySage', 2, 'YoungBufflalo', 25, False),
    ('Battle_TheBoySage', 3, 'YoungBufflalo', 25, False),
    ('Battle_TheBoySage', 5, 'TheBoySage', 30, True),
    
    ('Battle_BullDemonKing', 1, 'YoungBufflalo', 25, False),
    ('Battle_BullDemonKing', 2, 'BullDemonKing_Boss', 30, True),
    ('Battle_BullDemonKing', 3, 'YoungBufflalo', 25, False),
    ('Battle_BullDemonKing', 4, 'YoungBufflalo', 25, False),
    ('Battle_BullDemonKing', 5, 'TheBoySage', 30, False),
    ('Battle_BullDemonKing', 6, 'YoungBufflalo', 25, False),
]

saved = False
for attempt in range(10):
    try:
        wb = openpyxl.load_workbook('Tool/data/GameConfig.xlsx')
        ws = wb['StageEnemies']
        
        # Clear old rows except header
        while ws.max_row > 1:
            ws.delete_rows(2)
            
        for row in exact_levels:
            ws.append(list(row))
            
        wb.save('Tool/data/GameConfig.xlsx')
        saved = True
        print("Successfully saved all stage levels to GameConfig.xlsx!")
        break
    except PermissionError:
        print(f"Excel is currently locked, attempt {attempt+1}/10. Retrying in 1s...")
        time.sleep(1)

if not saved:
    print("WARNING: Could not save GameConfig.xlsx because file is locked by another program (Excel).")
else:
    # Verify contents
    wb = openpyxl.load_workbook('Tool/data/GameConfig.xlsx', data_only=True)
    ws = wb['StageEnemies']
    print(f"\nVerification: Total rows in StageEnemies = {ws.max_row}")
    for r in range(1, min(ws.max_row + 1, 45)):
        row_vals = [cell.value for cell in ws[r]]
        print(f"Row {r:2d}: {row_vals}")
