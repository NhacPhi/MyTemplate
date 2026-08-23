import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('Tool/data/GameConfig.xlsx', data_only=True)
ws = wb['SkillConfig']

header = [c.value for c in ws[1]]
print("Header:", header)

for r in list(ws.iter_rows(values_only=True))[1:15]:
    if r[0]:
        print(r)
