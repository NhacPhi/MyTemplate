import openpyxl, time, sys
sys.stdout.reconfigure(encoding='utf-8')

for attempt in range(10):
    try:
        wb = openpyxl.load_workbook('Tool/data/Localizations.xlsx')
        ws = wb['Effect'] if 'Effect' in wb.sheetnames else wb.active
        
        headers = [c.value for c in ws[1]]
        col_vn = headers.index('VIETNAMESE') if 'VIETNAMESE' in headers else 1
        col_en = headers.index('ENGLISH') if 'ENGLISH' in headers else 2
        
        for row in ws.iter_rows(min_row=2):
            k = str(row[0].value).strip() if row[0].value else ''
            if k == 'EFF_SILENCE_NAME':
                row[col_vn].value = 'Câm Lặng'
                row[col_en].value = 'Silence'
                print(f"Fixed EFF_SILENCE_NAME: VN='{row[col_vn].value}', EN='{row[col_en].value}'")
            elif k == 'EFF_SILENCE_DES':
                row[col_vn].value = 'Không thể sử dụng Tuyệt kỹ và Kỹ năng'
                row[col_en].value = 'Cannot use Major and Ultimate skills'
                print(f"Fixed EFF_SILENCE_DES: VN='{row[col_vn].value}', EN='{row[col_en].value}'")
                
        wb.save('Tool/data/Localizations.xlsx')
        print("Successfully updated Localizations.xlsx with correct column mappings!")
        break
    except PermissionError:
        print(f"Localizations.xlsx locked, attempt {attempt+1}, retrying in 1s...")
        time.sleep(1)
