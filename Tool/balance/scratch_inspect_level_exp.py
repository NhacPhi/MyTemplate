import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('Tool/data/GameConfig.xlsx', data_only=True)
print("Sheet names:", wb.sheetnames)

for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n--- Sheet: {name} (max_row={ws.max_row}, max_col={ws.max_column}) ---")
    headers = [cell.value for cell in ws[1]]
    print("Headers:", headers[:15])
    if ws.max_row > 1:
        first_row = [cell.value for cell in ws[2]]
        print("Row 2:", first_row[:15])
