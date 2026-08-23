import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('Tool/data/GameConfig.xlsx')
ws = wb['CharacterStat']

# Find header column indices
header = [cell.value for cell in ws[1]]
print("Header:", header)

id_idx = header.index('ID')
cr_idx = header.index('crit_rare') if 'crit_rare' in header else header.index('crit_rate')
res_idx = header.index('crit_dmg_res')

# Get Class from 'Character' sheet
ws_char = wb['Character']
char_class_map = {}
for r in ws_char.iter_rows(min_row=2, values_only=True):
    if r[0]:
        char_class_map[r[0]] = r[4]

for row in ws.iter_rows(min_row=2):
    cid = row[id_idx].value
    cls = char_class_map.get(cid, 'Character')
    if cls == 'Boss':
        row[cr_idx].value = 50
        row[res_idx].value = 50
        print(f"Updated Boss {cid}: crit_rate = 50%, crit_dmg_res = 50%")

wb.save('Tool/data/GameConfig.xlsx')
print("Successfully saved CharacterStat changes to GameConfig.xlsx")
