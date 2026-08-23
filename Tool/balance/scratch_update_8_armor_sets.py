import openpyxl
import time
import sys

sys.stdout.reconfigure(encoding='utf-8')

gc_path = 'Tool/data/GameConfig.xlsx'

sets_data = [
    ('Armor01', 'STR_ARMOR_01', 6, 'atk, crit_dmg', '20, 15', 'Percent, Percent'),
    ('Armor02', 'STR_ARMOR_02', 6, 'speed, hp', '8, 10', 'Percent, Percent'),
    ('Armor03', 'STR_ARMOR_03', 6, 'hp', '20', 'Percent'),
    ('Armor04', 'STR_ARMOR_04', 6, 'hp, def', '10, 10', 'Percent, Percent'),
    ('Armor05', 'STR_ARMOR_05', 6, 'atk, hp', '8, 8', 'Percent, Percent'),
    ('Armor06', 'STR_ARMOR_06', 6, 'atk, crit_dmg', '15, 15', 'Percent, Percent'),
    ('Armor07', 'STR_ARMOR_07', 6, 'crit_rate, crit_dmg', '10, 20', 'Percent, Percent'),
    ('Armor08', 'STR_ARMOR_08', 6, 'atk, penetration', '12, 10', 'Percent, Percent'),
]

pool_mapping = {
    'Armor01': 'Warrior_Pool',
    'Armor02': 'Warrior_Pool',
    'Armor03': 'Tanker_Pool',
    'Armor04': 'Tanker_Pool',
    'Armor05': 'Warrior_Pool',
    'Armor06': 'Assassin_Pool',
    'Armor07': 'Assassin_Pool',
    'Armor08': 'Warrior_Pool',
}

for attempt in range(5):
    try:
        wb = openpyxl.load_workbook(gc_path)

        # 1. Update SetBonusConfig sheet
        ws_set = wb['SetBonusConfig']
        ws_set.delete_rows(1, ws_set.max_row)
        ws_set.append(['ID', 'Name', 'Pieces_Required', 'Bonus_Stat_Type', 'Bonus_Value', 'Modifier_Type'])
        for row in sets_data:
            ws_set.append(list(row))

        # 2. Update Armor sheet substat_pool_id
        ws_armor = wb['Armor']
        for row in ws_armor.iter_rows(values_only=False):
            if row[0].value == 'ID': continue
            set_name = row[2].value
            if set_name in pool_mapping:
                row[7].value = pool_mapping[set_name]

        wb.save(gc_path)
        print("Successfully updated GameConfig.xlsx for all 8 Armor Sets!")
        break
    except Exception as e:
        print(f"Attempt {attempt+1} failed: {e}. Retrying...")
        time.sleep(1)
