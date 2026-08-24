import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('Tool/data/GameConfig.xlsx', data_only=True)
ws = wb['StageEnemies']

print("=== StageEnemies ===")
for r in range(1, ws.max_row + 1):
    row_vals = [cell.value for cell in ws[r]]
    if any(row_vals):
        print(f"Row {r:2d}: {row_vals}")
