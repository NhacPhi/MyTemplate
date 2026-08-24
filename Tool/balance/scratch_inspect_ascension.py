import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('Tool/data/GameConfig.xlsx', data_only=True)
ws = wb['AscensionConfig']

print("=== AscensionConfig ===")
for r in range(1, ws.max_row + 1):
    row_vals = [cell.value for cell in ws[r]]
    if any(row_vals):
        print(row_vals)
