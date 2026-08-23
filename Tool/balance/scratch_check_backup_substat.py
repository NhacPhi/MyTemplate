import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('Tool/data/backup/GameConfig.xlsx', data_only=True)
if 'SubstatPool' in wb.sheetnames:
    ws = wb['SubstatPool']
    for r in ws.iter_rows(values_only=True):
        print(r)
else:
    print("No SubstatPool sheet in backup")
