import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('Tool/data/Localizations.xlsx', data_only=True)
ws = wb['Effect']
print("Headers in sheet Effect:", [cell.value for cell in ws[1]])

for r in range(2, ws.max_row + 1):
    row_vals = [cell.value for cell in ws[r]]
    if any(row_vals):
        print(f"Row {r:2d}: {row_vals}")
